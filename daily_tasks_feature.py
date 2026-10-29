# -*- coding: utf-8 -*-
"""Daily Tasks feature using the same SQLAlchemy database as the main app.

Supports the application's configured SQLite or PostgreSQL engine. Attachment
metadata and uploaded bytes are stored in the same database so records do not
silently disappear when a deployment's local filesystem is ephemeral.
"""
from __future__ import annotations

import io
import os
import re
import uuid
from datetime import date, datetime
from functools import wraps
from pathlib import Path
from typing import Any
from collections.abc import Mapping

from sqlalchemy import (
    CheckConstraint, Column, Float, ForeignKey, Index, Integer, LargeBinary,
    MetaData, String, Table, Text, inspect, text,
)

from flask import (
    Blueprint, abort, current_app, flash, redirect, render_template, request,
    send_file, url_for,
)
from werkzeug.utils import secure_filename


daily_tasks_bp = Blueprint("daily_tasks", __name__, url_prefix="/daily-tasks")
ALLOWED_EXTENSIONS = {"png", "jpg", "jpeg", "webp", "pdf", "doc", "docx", "xls", "xlsx", "txt"}
MAX_ATTACHMENT_BYTES = 20 * 1024 * 1024
MAX_ATTACHMENTS_PER_TASK = 10


def _engine():
    """Return the primary SQLAlchemy engine configured by the main app."""
    extension = current_app.extensions.get("sqlalchemy")
    if extension is None:
        raise RuntimeError("Main SQLAlchemy extension is unavailable.")
    engine = extension.engines.get(None)
    if engine is None:
        raise RuntimeError("Could not resolve the main SQLAlchemy engine.")
    if engine.dialect.name not in {"sqlite", "postgresql"}:
        raise RuntimeError(
            f"Daily Tasks supports SQLite/PostgreSQL, not {engine.dialect.name}."
        )
    return engine


_metadata = MetaData()
_daily_tasks_table = Table(
    "daily_tasks",
    _metadata,
    Column("id", Integer, primary_key=True, autoincrement=True),
    Column("task_name", String(240), nullable=False),
    Column("work_area", String(200), nullable=False),
    Column("worker_count", Integer, nullable=False, default=0),
    Column("completion_pct", Float, nullable=False, default=0),
    Column("task_date", String(10), nullable=False),
    Column("notes", Text, nullable=False, default=""),
    Column("created_by", String(160), nullable=False, default=""),
    Column("created_at", String(32), nullable=False),
    Column("updated_at", String(32), nullable=False),
    CheckConstraint("worker_count >= 0", name="ck_daily_tasks_worker_count"),
    CheckConstraint(
        "completion_pct >= 0 AND completion_pct <= 100",
        name="ck_daily_tasks_completion_pct",
    ),
)
_daily_task_attachments_table = Table(
    "daily_task_attachments",
    _metadata,
    Column("id", Integer, primary_key=True, autoincrement=True),
    Column(
        "task_id", Integer,
        ForeignKey("daily_tasks.id", ondelete="CASCADE"),
        nullable=False,
    ),
    Column("original_name", String(255), nullable=False),
    Column("stored_name", String(80), nullable=False),
    Column("extension", String(10), nullable=False),
    Column("uploaded_at", String(32), nullable=False),
    # BYTEA on PostgreSQL; BLOB on SQLite. Existing records can remain NULL.
    Column("file_data", LargeBinary, nullable=True),
)
Index("idx_daily_tasks_date", _daily_tasks_table.c.task_date)
Index("idx_daily_tasks_area", _daily_tasks_table.c.work_area)


class _CompatRow(Mapping):
    """SQLAlchemy row mapping with sqlite.Row-like integer indexing."""
    def __init__(self, values):
        self._data = dict(values)
        self._keys = list(self._data.keys())
        self._values = list(self._data.values())

    def __getitem__(self, key):
        if isinstance(key, int):
            return self._values[key]
        return self._data[key]

    def __iter__(self):
        return iter(self._keys)

    def __len__(self):
        return len(self._keys)

    def keys(self):
        return self._data.keys()


class _CompatResult:
    def __init__(self, result, lastrowid=None):
        self._result = result
        self.lastrowid = lastrowid if lastrowid is not None else getattr(result, "lastrowid", None)
        self.rowcount = result.rowcount

    @staticmethod
    def _wrap(row):
        return _CompatRow(row._mapping) if row is not None else None

    def fetchone(self):
        return self._wrap(self._result.fetchone())

    def fetchall(self):
        return [self._wrap(row) for row in self._result.fetchall()]


class _CompatConnection:
    """Small bridge for legacy positional SQL, backed by SQLAlchemy transactions."""
    def __init__(self, engine):
        self._connection = engine.connect()
        self._engine = engine
        self._transaction = self._connection.begin()
        self._closed = False

    @staticmethod
    def _bind(statement, params):
        if params is None:
            return statement, {}
        if isinstance(params, Mapping):
            return statement, dict(params)
        values = list(params) if isinstance(params, (list, tuple)) else [params]
        index = 0

        def replace(_match):
            nonlocal index
            name = f"p{index}"
            index += 1
            return ":" + name

        converted = re.sub(r"\?", replace, statement)
        if index != len(values):
            raise ValueError(
                f"SQL placeholder count ({index}) does not match parameter count ({len(values)})."
            )
        return converted, {f"p{i}": value for i, value in enumerate(values)}

    def execute(self, statement, params=None):
        sql = str(statement).strip().rstrip(";")
        sql, bindings = self._bind(sql, params)
        returning_id = False
        if (
            self._engine.dialect.name == "postgresql"
            and re.match(r"(?is)^INSERT\s+INTO\s+\"?(daily_tasks|daily_task_attachments)\"?\b", sql)
            and not re.search(r"(?i)\bRETURNING\b", sql)
        ):
            sql += " RETURNING id"
            returning_id = True

        result = self._connection.execute(text(sql), bindings)
        if returning_id:
            row = result.fetchone()
            lastrowid = row[0] if row is not None else None
            return _CompatResult(result, lastrowid=lastrowid)
        return _CompatResult(result)

    def commit(self):
        if self._transaction is not None and self._transaction.is_active:
            self._transaction.commit()
        if not self._closed:
            self._transaction = self._connection.begin()

    def rollback(self):
        if self._transaction is not None and self._transaction.is_active:
            self._transaction.rollback()
        if not self._closed:
            self._transaction = self._connection.begin()

    def close(self):
        if self._closed:
            return
        try:
            if self._transaction is not None and self._transaction.is_active:
                self._transaction.rollback()
        finally:
            self._connection.close()
            self._closed = True

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc_value, traceback):
        try:
            if self._transaction is not None and self._transaction.is_active:
                if exc_type is None:
                    self._transaction.commit()
                else:
                    self._transaction.rollback()
        finally:
            self._connection.close()
            self._closed = True
        return False


def _uploads_path() -> str:
    path = os.path.join(
        current_app.static_folder or os.path.join(current_app.root_path, "static"),
        "uploads", "daily_tasks",
    )
    os.makedirs(path, exist_ok=True)
    return path


def _connect() -> _CompatConnection:
    return _CompatConnection(_engine())


def _init_db() -> None:
    """Create feature tables in the application's configured primary database."""
    engine = _engine()
    _metadata.create_all(engine)

    # create_all does not add columns to a table that already exists.
    column_names = {
        item["name"]
        for item in inspect(engine).get_columns("daily_task_attachments")
    }
    if "file_data" not in column_names:
        if engine.dialect.name == "sqlite":
            ddl_type = "BLOB"
        elif engine.dialect.name == "postgresql":
            ddl_type = "BYTEA"
        else:
            raise RuntimeError("Unsupported database type for attachment storage migration.")
        with engine.begin() as connection:
            connection.execute(
                text(f"ALTER TABLE daily_task_attachments ADD COLUMN file_data {ddl_type}")
            )

def _ensure_db(fn):
    @wraps(fn)
    def wrapper(*args, **kwargs):
        _init_db()
        return fn(*args, **kwargs)
    return wrapper


def _date_filter() -> tuple[str, str]:
    start = (request.values.get("start_date") or "").strip()
    end = (request.values.get("end_date") or "").strip()
    for value in (start, end):
        if value:
            try:
                datetime.strptime(value, "%Y-%m-%d")
            except ValueError:
                abort(400, description="صيغة التاريخ غير صحيحة. استخدم YYYY-MM-DD.")
    if start and end and start > end:
        abort(400, description="تاريخ البداية يجب ألا يتجاوز تاريخ النهاية.")
    return start, end


def _fetch_tasks(start: str = "", end: str = "") -> list[dict[str, Any]]:
    sql = """
        SELECT t.*,
               (SELECT COUNT(*) FROM daily_task_attachments a WHERE a.task_id=t.id) AS attachment_count
        FROM daily_tasks t
        WHERE 1=1
    """
    params: list[Any] = []
    if start:
        sql += " AND t.task_date >= ?"
        params.append(start)
    if end:
        sql += " AND t.task_date <= ?"
        params.append(end)
    sql += " ORDER BY t.task_date DESC, t.id DESC"
    with _connect() as con:
        return [dict(row) for row in con.execute(sql, params).fetchall()]


def _actor() -> str:
    # Do not assume a particular user model; use the common session identity keys if present.
    from flask import session
    for key in ("username", "user_name", "email", "user_id", "user"):
        value = session.get(key)
        if value is not None and not isinstance(value, (dict, list)):
            return str(value)[:160]
    try:
        from flask_login import current_user
        if getattr(current_user, "is_authenticated", False):
            return str(getattr(current_user, "username", None) or getattr(current_user, "email", None) or getattr(current_user, "id", ""))[:160]
    except Exception:
        pass
    return ""


def _allowed_file(filename: str) -> bool:
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


def _save_attachments(task_id: int) -> None:
    files = [f for f in request.files.getlist("attachments") if f and f.filename]
    if len(files) > MAX_ATTACHMENTS_PER_TASK:
        raise ValueError(f"الحد الأقصى للمرفقات في المهمة الواحدة هو {MAX_ATTACHMENTS_PER_TASK} ملفات.")
    upload_dir = _uploads_path()
    saved: list[str] = []
    with _connect() as con:
        try:
            for file in files:
                # Preserve Arabic display names; only the random stored name is used on disk.
                original = os.path.basename((file.filename or "").replace("\\", "/")).replace("\x00", "").strip()[:255]
                ext = Path(original).suffix.lower().lstrip(".")
                if not original or ext not in ALLOWED_EXTENSIONS:
                    raise ValueError(f"نوع الملف غير مسموح: {file.filename}. الأنواع المتاحة: صور، PDF، Word، Excel، TXT.")
                data = file.read(MAX_ATTACHMENT_BYTES + 1)
                if len(data) > MAX_ATTACHMENT_BYTES:
                    raise ValueError(f"حجم الملف يتجاوز 20 ميغابايت: {original}")
                if not data:
                    continue
                stored = f"{uuid.uuid4().hex}.{ext}"
                destination = os.path.join(upload_dir, stored)
                with open(destination, "wb") as out:
                    out.write(data)
                saved.append(destination)
                con.execute(
                    "INSERT INTO daily_task_attachments(task_id, original_name, stored_name, extension, uploaded_at, file_data) VALUES (?, ?, ?, ?, ?, ?)",
                    (task_id, original, stored, ext, datetime.now().isoformat(timespec="seconds"), data),
                )
            con.commit()
        except Exception:
            con.rollback()
            for path in saved:
                try:
                    os.remove(path)
                except OSError:
                    pass
            raise


def _set_task_fields(task_id: int | None = None) -> int:
    task_name = (request.form.get("task_name") or "").strip()
    work_area = (request.form.get("work_area") or "").strip()
    task_date = (request.form.get("task_date") or date.today().isoformat()).strip()
    notes = (request.form.get("notes") or "").strip()[:4000]
    try:
        workers = int(request.form.get("worker_count", "0"))
        completion = float(request.form.get("completion_pct", "0"))
        datetime.strptime(task_date, "%Y-%m-%d")
    except (ValueError, TypeError):
        raise ValueError("يرجى إدخال عدد العمال ونسبة الإنجاز والتاريخ بشكل صحيح.")
    if not task_name or not work_area:
        raise ValueError("حقلا مهمة العمل ومنطقة العمل مطلوبان.")
    if workers < 0 or workers > 100000:
        raise ValueError("عدد العمال يجب أن يكون رقمًا صحيحًا غير سالب.")
    if not 0 <= completion <= 100:
        raise ValueError("نسبة الإنجاز يجب أن تكون بين 0 و100.")
    now = datetime.now().isoformat(timespec="seconds")
    with _connect() as con:
        if task_id is None:
            cur = con.execute(
                "INSERT INTO daily_tasks(task_name, work_area, worker_count, completion_pct, task_date, notes, created_by, created_at, updated_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                (task_name, work_area, workers, completion, task_date, notes, _actor(), now, now),
            )
            task_id = int(cur.lastrowid)
        else:
            exists = con.execute("SELECT id FROM daily_tasks WHERE id=?", (task_id,)).fetchone()
            if not exists:
                abort(404)
            con.execute(
                "UPDATE daily_tasks SET task_name=?, work_area=?, worker_count=?, completion_pct=?, task_date=?, notes=?, updated_at=? WHERE id=?",
                (task_name, work_area, workers, completion, task_date, notes, now, task_id),
            )
        con.commit()
    return task_id


@daily_tasks_bp.route("/", methods=["GET", "POST"], endpoint="index")
@_ensure_db
def index():
    if request.method == "POST":
        try:
            task_id = _set_task_fields()
            _save_attachments(task_id)
            flash("تمت إضافة المهمة اليومية بنجاح.", "success")
            return redirect(url_for("daily_tasks.index", embed=request.args.get("embed") or None))
        except ValueError as exc:
            flash(str(exc), "danger")
        except Exception as exc:
            current_app.logger.exception("daily task create failed")
            flash("تعذر حفظ المهمة أو المرفقات. تحقق من الملفات ثم حاول مجددًا.", "danger")
    start, end = _date_filter()
    all_matching_tasks = _fetch_tasks(start, end)
    filtered_task_count = len(all_matching_tasks)
    page_size = 10
    try:
        page = max(1, int(request.args.get("page", "1")))
    except (TypeError, ValueError):
        page = 1
    pages = max(1, (filtered_task_count + page_size - 1) // page_size)
    page = min(page, pages)
    offset = (page - 1) * page_size
    tasks = all_matching_tasks[offset:offset + page_size]
    with _connect() as con:
        total = con.execute("SELECT COUNT(*) FROM daily_tasks").fetchone()[0]
        avg_completion = con.execute("SELECT COALESCE(AVG(completion_pct),0) FROM daily_tasks").fetchone()[0]
        total_workers = con.execute("SELECT COALESCE(SUM(worker_count),0) FROM daily_tasks").fetchone()[0]
    return render_template(
        "daily_tasks/index.html", tasks=tasks, start_date=start, end_date=end,
        today=date.today().isoformat(), total_tasks=total,
        page=page, pages=pages, page_size=page_size, filtered_task_count=filtered_task_count,
        avg_completion=round(float(avg_completion or 0), 1), total_workers=total_workers,
    )


@daily_tasks_bp.route("/<int:task_id>/edit", methods=["GET", "POST"], endpoint="edit")
@_ensure_db
def edit(task_id: int):
    with _connect() as con:
        row = con.execute("SELECT * FROM daily_tasks WHERE id=?", (task_id,)).fetchone()
        attachments = [dict(x) for x in con.execute("SELECT * FROM daily_task_attachments WHERE task_id=? ORDER BY id", (task_id,)).fetchall()]
    if not row:
        abort(404)
    if request.method == "POST":
        try:
            _set_task_fields(task_id)
            _save_attachments(task_id)
            flash("تم تحديث المهمة بنجاح.", "success")
            return redirect(url_for("daily_tasks.index", embed=request.args.get("embed") or None))
        except ValueError as exc:
            flash(str(exc), "danger")
        except Exception:
            current_app.logger.exception("daily task edit failed")
            flash("تعذر تحديث المهمة أو حفظ المرفقات.", "danger")
    return render_template("daily_tasks/edit.html", task=dict(row), attachments=attachments)


@daily_tasks_bp.route("/<int:task_id>/delete", methods=["POST"], endpoint="delete")
@_ensure_db
def delete(task_id: int):
    with _connect() as con:
        rows = con.execute(
            "SELECT stored_name, (file_data IS NULL) AS is_legacy FROM daily_task_attachments WHERE task_id=?",
            (task_id,),
        ).fetchall()
        con.execute("DELETE FROM daily_task_attachments WHERE task_id=?", (task_id,))
        deleted = con.execute("DELETE FROM daily_tasks WHERE id=?", (task_id,)).rowcount
        con.commit()
    # Old attachments may exist only as files. New attachments are backed by the database.
    for row in rows:
        if row["is_legacy"]:
            try:
                os.remove(os.path.join(_uploads_path(), row["stored_name"]))
            except OSError:
                pass
    flash("تم حذف المهمة." if deleted else "المهمة غير موجودة.", "success" if deleted else "warning")
    return redirect(url_for("daily_tasks.index", embed=request.args.get("embed") or None))


@daily_tasks_bp.route("/attachment/<int:attachment_id>", endpoint="attachment")
@_ensure_db
def attachment(attachment_id: int):
    with _connect() as con:
        row = con.execute("SELECT * FROM daily_task_attachments WHERE id=?", (attachment_id,)).fetchone()
    if not row:
        abort(404)
    payload = row["file_data"]
    if payload is not None:
        return send_file(
            io.BytesIO(payload),
            mimetype="application/octet-stream",
            as_attachment=True,
            download_name=row["original_name"],
            max_age=0,
        )
    # Backwards compatibility for attachments saved before database-backed storage.
    path = os.path.join(_uploads_path(), row["stored_name"])
    if not os.path.isfile(path):
        abort(404)
    return send_file(path, as_attachment=True, download_name=row["original_name"])


@daily_tasks_bp.route("/export.xlsx", endpoint="export_xlsx")
@_ensure_db
def export_xlsx():
    start, end = _date_filter()
    tasks = _fetch_tasks(start, end)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError:
        flash("تصدير Excel يتطلب تثبيت المكتبة openpyxl داخل بيئة التطبيق.", "danger")
        return redirect(url_for("daily_tasks.index", start_date=start, end_date=end))
    wb = Workbook()
    ws = wb.active
    ws.title = "المهام اليومية"
    ws.sheet_view.rightToLeft = True
    ws.merge_cells("A1:H1")
    ws["A1"] = "تقرير المهام اليومية"
    ws["A1"].font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="17365D")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30
    ws.merge_cells("A2:H2")
    ws["A2"] = f"الفترة: {start or 'من البداية'} إلى {end or 'حتى اليوم'} | تاريخ التصدير: {date.today().isoformat()}"
    ws["A2"].font = Font(name="Arial", size=10, italic=True, color="404040")
    ws["A2"].alignment = Alignment(horizontal="right")
    headers = ["م", "التاريخ", "مهمة العمل", "منطقة العمل", "عدد العمال", "نسبة الإنجاز %", "عدد المرفقات", "ملاحظات"]
    header_row = 4
    ws.append([""] * 8)
    ws.append(headers)
    for cell in ws[header_row]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, task in enumerate(tasks, 1):
        ws.append([i, task["task_date"], task["task_name"], task["work_area"], task["worker_count"], float(task["completion_pct"]), task["attachment_count"], task["notes"]])
    widths = [7, 15, 34, 25, 14, 17, 14, 36]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + idx)].width = width
    thin = Side(style="thin", color="D9E2F3")
    for row in ws.iter_rows(min_row=header_row, max_row=ws.max_row, min_col=1, max_col=8):
        for cell in row:
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(horizontal="right" if cell.column in (3, 4, 8) else "center", vertical="center", wrap_text=True)
            if cell.row > header_row:
                cell.font = Font(name="Arial", size=10)
    if ws.max_row > header_row:
        table = Table(displayName="DailyTasksTable", ref=f"A{header_row}:H{ws.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
        ws.add_table(table)
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A{header_row}:H{max(header_row, ws.max_row)}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:4"
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    suffix = f"_{start or 'all'}_{end or date.today().isoformat()}"
    return send_file(buffer, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", as_attachment=True, download_name=f"daily_tasks{suffix}.xlsx")


@daily_tasks_bp.route("/export.pdf", endpoint="export_pdf")
@_ensure_db
def export_pdf():
    start, end = _date_filter()
    tasks = _fetch_tasks(start, end)

    try:
        import arabic_reshaper
        from bidi.algorithm import get_display
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.enums import TA_RIGHT, TA_CENTER
        from reportlab.lib.units import mm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import SimpleDocTemplate, Paragraph, LongTable, TableStyle
        from xml.sax.saxutils import escape
    except ImportError as exc:
        current_app.logger.exception("Daily Tasks PDF dependencies are unavailable")
        abort(500, description="PDF export dependencies are missing: " + str(exc))

    def rtl(value):
        raw = "" if value is None else str(value)
        try:
            return get_display(arabic_reshaper.reshape(raw))
        except Exception:
            return raw

    # Find a Unicode font that supports Arabic.
    fonts_dir = os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts")
    candidates = [
        os.path.join(fonts_dir, n)
        for n in ("arial.ttf", "tahoma.ttf", "segoeui.ttf", "calibri.ttf")
    ]
    candidates += [
        "/usr/share/fonts/truetype/noto/NotoSansArabic-Regular.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
    ]

    if os.path.isdir(fonts_dir):
        for pattern in ("arial*.ttf", "tahoma*.ttf", "segoeui*.ttf", "calibri*.ttf"):
            candidates.extend(str(p) for p in Path(fonts_dir).glob(pattern))

    font_name = None
    for font_path in candidates:
        if not os.path.isfile(font_path):
            continue
        try:
            if "DailyTasksUnicode" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("DailyTasksUnicode", font_path))
            pdfmetrics.stringWidth(rtl("المهام اليومية 123 ABC"), "DailyTasksUnicode", 9)
            font_name = "DailyTasksUnicode"
            break
        except Exception:
            continue

    if not font_name:
        current_app.logger.error("No usable Unicode font found for Daily Tasks PDF")
        abort(
            500,
            description="لم يتم العثور على خط يدعم العربية. تحقق من وجود Arial أو Tahoma في مجلد خطوط Windows."
        )

    def para(value, style):
        text = escape(rtl(value)).replace("\r\n", "\n").replace("\n", "<br/>")
        return Paragraph(text or " ", style)

    page_size = landscape(A4)
    page_w, page_h = page_size
    margin = 10 * mm
    usable_w = page_w - 2 * margin

    head_style = ParagraphStyle(
        "DTPdfHead", fontName=font_name, fontSize=8, leading=10,
        alignment=TA_CENTER, textColor=colors.white, wordWrap="CJK"
    )
    cell_style = ParagraphStyle(
        "DTPdfCell", fontName=font_name, fontSize=7.5, leading=9.5,
        alignment=TA_RIGHT, textColor=colors.HexColor("#1f2937"), wordWrap="CJK"
    )
    center_style = ParagraphStyle(
        "DTPdfCenter", parent=cell_style, alignment=TA_CENTER
    )

    raw_widths = [28, 62, 165, 120, 58, 68, 55, 242]
    factor = usable_w / sum(raw_widths)
    widths = [n * factor for n in raw_widths]

    headers = [
        "م", "التاريخ", "مهمة العمل", "منطقة العمل",
        "عدد العمال", "نسبة الإنجاز", "المرفقات", "ملاحظات"
    ]
    data = [[para(v, head_style) for v in headers]]

    for i, task in enumerate(tasks, 1):
        data.append([
            para(i, center_style),
            para(task.get("task_date", ""), center_style),
            para(task.get("task_name", ""), cell_style),
            para(task.get("work_area", ""), cell_style),
            para(task.get("worker_count", 0), center_style),
            para(f"{float(task.get('completion_pct', 0) or 0):g}%", center_style),
            para(task.get("attachment_count", 0), center_style),
            para(task.get("notes", ""), cell_style),
        ])

    if not tasks:
        data.append([para("لا توجد مهام ضمن الفترة المحددة.", cell_style)] + [""] * 7)
    # ReportLab paints columns from left to right; reverse cells and widths so
    # readers scanning Arabic from the right see the intended logical order.
    data = [list(reversed(row)) for row in data]
    widths = list(reversed(widths))

    table = LongTable(
        data, colWidths=widths, repeatRows=1, hAlign="LEFT", splitByRow=1
    )
    commands = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#17365D")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, 0), 7),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 7),
        ("TOPPADDING", (0, 1), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 1), (-1, -1), 5),
        ("ALIGN", (0, 0), (-1, -1), "RIGHT"),
        ("ALIGN", (1, 0), (3, -1), "CENTER"),
        ("ALIGN", (6, 0), (7, -1), "CENTER"),
    ]
    for row in range(2, len(data), 2):
        commands.append(
            ("BACKGROUND", (0, row), (-1, row), colors.HexColor("#F2F6FC"))
        )
    table.setStyle(TableStyle(commands))

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=page_size,
        rightMargin=margin,
        leftMargin=margin,
        topMargin=27 * mm,
        bottomMargin=15 * mm,
        title="Daily Tasks Report",
        author="Bin Hayell"
    )

    title_text = rtl("تقرير المهام اليومية")
    period_text = rtl(f"الفترة: {start or 'من البداية'} — {end or 'حتى اليوم'}")
    footer_text = rtl(
        f"تاريخ التصدير: {date.today().isoformat()} | إجمالي المهام: {len(tasks)}"
    )

    def draw_page(canv, document):
        canv.saveState()
        canv.setFillColor(colors.HexColor("#17365D"))
        canv.rect(0, page_h - 22 * mm, page_w, 22 * mm, fill=1, stroke=0)
        canv.setFillColor(colors.white)
        canv.setFont(font_name, 14)
        canv.drawRightString(page_w - margin, page_h - 9 * mm, title_text)
        canv.setFont(font_name, 8)
        canv.drawRightString(page_w - margin, page_h - 16 * mm, period_text)
        canv.setFillColor(colors.HexColor("#64748b"))
        canv.setFont(font_name, 7.5)
        canv.drawRightString(page_w - margin, 7 * mm, footer_text)
        canv.drawString(margin, 7 * mm, rtl(f"صفحة {document.page}"))
        canv.restoreState()

    doc.build([table], onFirstPage=draw_page, onLaterPages=draw_page)
    pdf_bytes = buffer.getvalue()

    if not pdf_bytes.startswith(b"%PDF-") or len(pdf_bytes) < 700:
        current_app.logger.error("PDF output invalid/empty; bytes=%s", len(pdf_bytes))
        abort(500, description="تعذر إنشاء ملف PDF صالح. راجع سجل الخادم.")

    suffix = f"_{start or 'all'}_{end or date.today().isoformat()}"
    response = send_file(
        io.BytesIO(pdf_bytes),
        mimetype="application/pdf",
        as_attachment=True,
        download_name=f"daily_tasks{suffix}.pdf",
        max_age=0
    )
    response.headers["X-Daily-Tasks-PDF-Bytes"] = str(len(pdf_bytes))
    response.headers["Cache-Control"] = "no-store"
    return response

def _daily_tasks_auth_guard(view):
    """Require a Flask-Login session without relying on a potentially missing auth.login endpoint.

    Some Bin Hayell builds expose the login form through the React frontend and API,
    while the app's global unauthorized callback points at the unregistered endpoint
    ``auth.login``. Flask-Login's stock decorator then produces HTTP 500 for anonymous
    requests. This feature-local guard preserves authentication and redirects anonymous
    visitors to a frontend login path instead of changing global authentication behavior.
    """
    @wraps(view)
    def guarded(*args, **kwargs):
        try:
            from flask_login import current_user
            authenticated = bool(current_user.is_authenticated)
        except Exception:
            authenticated = False
        if not authenticated:
            from urllib.parse import urlencode
            from flask import current_app, redirect, request, url_for
            target = None
            manager = getattr(current_app, "login_manager", None)
            login_view = getattr(manager, "login_view", None) if manager else None
            if login_view and login_view in current_app.view_functions:
                try:
                    target = url_for(login_view)
                except Exception:
                    target = None
            if target is None and "auth.login" in current_app.view_functions:
                try:
                    target = url_for("auth.login")
                except Exception:
                    target = None
            if target is None:
                target = str(current_app.config.get("DAILY_TASKS_LOGIN_URL") or "/login")
            separator = "&" if "?" in target else "?"
            next_url = request.full_path if request.query_string else request.path
            return redirect(target + separator + urlencode({"next": next_url}))
        return view(*args, **kwargs)
    setattr(guarded, "_daily_tasks_auth_wrapped", True)
    return guarded


def register_daily_tasks(app, login_required=None):
    """Register the feature with a feature-local authentication guard."""
    if "daily_tasks" not in app.blueprints:
        app.register_blueprint(daily_tasks_bp)

    # Use a safe local guard for this add-on. Do not rely on the app's global
    # unauthorized callback: some builds point it at a missing auth.login endpoint.
    for endpoint, view in list(app.view_functions.items()):
        if endpoint.startswith("daily_tasks.") and not getattr(view, "_daily_tasks_auth_wrapped", False):
            app.view_functions[endpoint] = _daily_tasks_auth_guard(view)

    # Create Daily Tasks tables in the same configured SQLite/PostgreSQL database as the app.
    with app.app_context():
        _init_db()
